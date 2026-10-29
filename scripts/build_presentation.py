"""Générer le support de soutenance HRNova et ses aperçus sans dépendance à l'API.

uv run --no-project --with python-pptx==1.0.2 python scripts/build_presentation.py
"""

import argparse
import json
import math
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.chart.data import ChartData
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
W, H, PPI = 13.333333, 7.5, 120
NAVY, TEAL, LIME = "102C3A", "007F82", "C7EA73"
BG, WHITE, INK, MUTED, LINE = "F4F7F8", "FFFFFF", "183744", "526A75", "DCE5E8"
ORANGE, PALE = "BA5B24", "E3F2EE"


def rgb(value):
    return RGBColor.from_string(value)


def font_file(bold=False):
    names = (
        ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"]
        if bold else
        ["/System/Library/Fonts/Supplemental/Arial.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
    )
    for name in names:
        if Path(name).is_file():
            return name
    return "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"


class Deck:
    def __init__(self, total=20):
        self.total = total
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.images, self.notes, self.slide = [], [], None
        self.prs.core_properties.title = "HRNova — Du modèle à une API traçable"
        self.prs.core_properties.author = "Samuel Gabbai"
        self.prs.core_properties.subject = "Soutenance Data Science : modèle, API, PostgreSQL, traçabilité et CI/CD"
        self.prs.core_properties.created = datetime(2026, 10, 27)

    def new(self, title, section, note, source="", dark=False):
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(NAVY if dark else BG)
        self.canvas = Image.new("RGB", (1600, 900), "#" + (NAVY if dark else BG))
        self.draw = ImageDraw.Draw(self.canvas)
        self.dark = dark
        self.box(0, 0, W, 0.075, TEAL)
        self.text(section.upper(), .55, .29, 12.1, .26, 10.5, LIME if dark else TEAL, bold=True)
        self.text(title, .55, .76, 12.1, .94, 28, WHITE if dark else NAVY, bold=True)
        self.text("SAMUEL GABBAI  /  HRNOVA", .55, 7.08, 5, .2, 9, "BDD1D6" if dark else MUTED)
        self.text(f"{len(self.prs.slides):02d} / {self.total}", 11.65, 7.06, 1.12, .23, 10, "BDD1D6" if dark else MUTED, align="right")
        if source:
            self.text(source, .55, 6.77, 12.1, .22, 8.5, "BDD1D6" if dark else MUTED)
        self.slide.notes_slide.notes_text_frame.text = note + ("\n\nSources : " + source if source else "")
        self.notes.append((title, note, source))

    def box(self, x, y, w, h, fill=WHITE, border=None, rounded=False):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if border:
            shape.line.color.rgb = rgb(border)
            shape.line.width = Pt(.6)
        else:
            shape.line.fill.background()
        rect = tuple(round(v * PPI) for v in (x, y, x + w, y + h))
        if rounded:
            shape.adjustments[0] = .10
            self.draw.rounded_rectangle(rect, radius=12, fill="#" + fill, outline="#" + border if border else None)
        else:
            self.draw.rectangle(rect, fill="#" + fill, outline="#" + border if border else None)
        return shape

    def text(self, content, x, y, w, h, size=18, color=INK, bold=False, align="left"):
        content = str(content)
        font = ImageFont.truetype(font_file(bold), max(1, round(size * PPI / 72)))
        lines = []
        for raw in content.split("\n"):
            words, line = raw.split(), ""
            if not words:
                lines.append("")
                continue
            for word in words:
                candidate = (line + " " + word).strip()
                if font.getlength(candidate) > w * PPI - 3 and line:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            lines.append(line)
        line_height = size * 1.18 / 72
        if len(lines) * line_height > h + .035:
            raise ValueError(f"Texte trop haut, slide {len(self.prs.slides)} : {content[:90]} ({len(lines)} lignes dans {h})")
        if any(font.getlength(line) > w * PPI + 4 for line in lines):
            raise ValueError(f"Mot trop large : {content[:90]}")
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold = "Arial", Pt(size), bold
            p.font.color.rgb = rgb(color)
            p.line_spacing = Pt(size * 1.18)
            p.space_before = p.space_after = Pt(0)
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            px = x * PPI
            if align == "center":
                px += (w * PPI - font.getlength(line)) / 2
            elif align == "right":
                px += w * PPI - font.getlength(line)
            self.draw.text((round(px), round(y * PPI + i * line_height * PPI)), line, font=font, fill="#" + color)
        return shape

    def arrow(self, x1, y1, x2, y2, color=TEAL, width=2):
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        conn.line.color.rgb, conn.line.width = rgb(color), Pt(width)
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        conn.line._get_or_add_ln().append(end)
        self.draw.line((round(x1 * PPI), round(y1 * PPI), round(x2 * PPI), round(y2 * PPI)), fill="#" + color, width=round(width * PPI / 72))
        angle = math.atan2(y2-y1, x2-x1)
        tip = (x2 * PPI, y2 * PPI)
        points = [tip] + [(tip[0]-12*math.cos(angle+a), tip[1]-12*math.sin(angle+a)) for a in (-.4, .4)]
        self.draw.polygon(points, fill="#" + color)

    def card(self, x, y, w, h, title, body, fill=WHITE, accent=TEAL, size=18):
        self.box(x, y, w, h, fill, rounded=True)
        self.box(x, y, .045, h, accent)
        self.text(title, x+.23, y+.24, w-.46, .60, 19, accent, bold=True)
        self.text(body, x+.23, y+1.02, w-.46, h-1.19, size)

    def stat(self, x, y, w, value, label, dark=False):
        self.text(value, x, y, w, .70, 38, LIME if dark else TEAL, bold=True)
        self.text(label, x, y+.78, w, .70, 17, WHITE if dark else INK)

    def banner(self, body, y=6.0, color=TEAL, fill=PALE):
        self.box(.55, y, 12.2, .55, fill, rounded=True)
        self.text(body, .76, y+.12, 11.78, .32, 14, color, bold=True)

    def table(self, headers, rows, x, y, widths, row_h=.65, font=16):
        self.box(x, y, sum(widths), row_h, NAVY)
        offset=x
        for title, w in zip(headers, widths):
            self.text(title, offset+.14, y+.14, w-.28, row_h-.17, font, WHITE, bold=True)
            offset+=w
        for i, row in enumerate(rows):
            cy=y+(i+1)*row_h
            self.box(x, cy, sum(widths), row_h, WHITE if i%2==0 else "EAF0F2")
            offset=x
            for value, w in zip(row, widths):
                self.text(value, offset+.14, cy+.13, w-.28, row_h-.16, font)
                offset+=w

    def finish(self):
        self.images.append(self.canvas)

    def save(self, out):
        out.mkdir(parents=True, exist_ok=True)
        target=out/"HRNova_Soutenance_2026-10-27.pptx"
        self.prs.save(target)
        thumb_w, thumb_h=400,225
        overview=Image.new("RGB",(thumb_w*4,thumb_h*5),"white")
        for i,img in enumerate(self.images):
            overview.paste(img.resize((thumb_w,thumb_h)),((i%4)*thumb_w,(i//4)*thumb_h))
        overview.save(out/"HRNova_Apercu.png")
        preview_dir=out/"apercus"
        preview_dir.mkdir(exist_ok=True)
        for i in (0,5,7,8,9,11,14,16):
            self.images[i].save(preview_dir/f"slide_{i+1:02d}.png")
        notes=["# HRNova — Notes pour la soutenance", "", "20 diapositives : 18 principales et 2 annexes. Les notes sont aussi intégrées au PowerPoint.", ""]
        for i,(title,note,source) in enumerate(self.notes,1):
            notes += [f"## {i:02d} — {title}", "", note, "", f"Source : {source}", ""]
        (out/"HRNova_Notes_orateur.md").write_text("\n".join(notes),encoding="utf-8")
        print(f"Créé : {target} ({len(self.prs.slides)} diapositives, notes intégrées).")


def build(out):
    metrics=json.loads((ROOT/"docs/model_metrics.json").read_text())
    native=metrics["reported_test_metrics_rounded"]["native_threshold_0_5"]
    tuned=metrics["reported_test_metrics_rounded"]["optimized_threshold"]
    d=Deck()

    d.new("Du modèle à une API traçable", "HRNova / TechNova Partners", "Présenter le projet comme un parcours complet : partir de trois sources RH, estimer une probabilité de départ puis rendre le calcul utilisable et vérifiable. Le périmètre actuel est un service local, avec publication du modèle sur le Hub. Les performances viennent du notebook enregistré et les tests logiciels vérifient le fonctionnement du service.", "État du projet examiné le 27 octobre 2026", dark=True)
    d.text("Anticiper les départs.\nComprendre chaque résultat.\nRetrouver chaque échange.", .65, 2.0, 8.9, 2.15, 30, WHITE, bold=True)
    d.box(10.25, 2.0, 2.45, 2.4, TEAL, rounded=True)
    d.text("DATA\nMODEL\nAPI",10.57,2.33,1.84,1.8,30,WHITE,bold=True)
    d.stat(.65, 5.10, 3.6, "1 470", "salariés", True)
    d.stat(4.85, 5.10, 3.6, "43", "variables préparées", True)
    d.stat(9.05, 5.10, 3.6, "7", "endpoints documentés", True)
    d.finish()

    d.new("Un outil d'aide à l'analyse RH", "01 / Besoin métier", "L'objectif est de mieux comprendre et repérer les situations associées aux départs. La cible connue est binaire : salarié resté ou parti. Le résultat est une probabilité et une décision selon un seuil. Souligner que le modèle ne prouve pas les causes de départ et ne remplace pas une analyse métier. Il faut relier les fausses alertes et les départs manqués à des coûts et à des actions définies avec les RH.", "Sources : README.md ; docs/model.md")
    d.card(.55,2.05,3.93,3.5,"COMPRENDRE","Relier données RH, évaluations et sondages.\n\nIdentifier des associations statistiques.")
    d.card(4.65,2.05,3.93,3.5,"ESTIMER","Produire une probabilité de départ et une décision au seuil choisi.")
    d.card(8.75,2.05,3.99,3.5,"TRACER","Retrouver les données utilisées, la version du modèle et le résultat.")
    d.banner("La qualité métier se mesure sur des données indépendantes, au-delà du bon fonctionnement de l'API.")
    d.finish()

    d.new("Trois sources, une population de 1 470 salariés", "02 / Données", "Les trois fichiers ont chacun 1 470 lignes, mais concernent la même population : on ne dispose pas de 4 410 salariés différents. Les clés sont rapprochées entre les sources, tout en conservant leur représentation d'origine en base. La classe positive représente 237 départs sur 1 470, soit environ 16,1 %. Le déséquilibre justifie de regarder précision, rappel et F1, pas uniquement l'accuracy.", "Sources : Data/*.csv ; docs/model.md ; docs/database.md")
    for x,title,body in [( .55,"SIRH","Âge, rémunération, poste et ancienneté"),(4.65,"SONDAGE","Engagement, contexte et cible de départ"),(8.75,"ÉVALUATIONS","Satisfaction, notes et heures supplémentaires")]:
        d.card(x,1.87,3.95,1.95,title,body,size=16)
    data=ChartData();data.categories=["Restés","Partis"];data.add_series("Salariés",[1233,237])
    chart=d.slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,Inches(.9),Inches(4.03),Inches(2.6),Inches(2.35),data).chart
    chart.has_legend=False; chart.has_title=False; chart.plots[0].hole_size=74
    for point,color in zip(chart.series[0].points,[LINE,TEAL]):
        point.format.fill.solid(); point.format.fill.fore_color.rgb=rgb(color); point.format.line.fill.background()
    bounds=(round(1.06*PPI),round(4.05*PPI),round(3.3*PPI),round(6.29*PPI))
    d.draw.arc(bounds,0,360,fill="#"+LINE,width=30)
    d.draw.arc(bounds,-90,-90+360*237/1470,fill="#"+TEAL,width=30)
    d.text("16,1 %",1.39,4.92,1.58,.52,25,TEAL,bold=True,align="center")
    d.text("de départs",1.40,5.48,1.55,.28,12,MUTED,align="center")
    d.stat(4.15,4.20,3.75,"237","départs • classe 1")
    d.stat(8.40,4.20,3.8,"1 233","salariés restés • classe 0")
    d.finish()

    d.new("Transformer les données en 43 variables numériques", "03 / Préparation", "La préparation est exécutée dans le notebook. Elle fusionne les sources, convertit les pourcentages, encode les variables et crée notamment des indicateurs de rémunération relative et une satisfaction moyenne. La cible et les identifiants sont exclus des variables explicatives. L'API reçoit le résultat de cette préparation ; elle ne fait pas la transformation d'un fichier RH brut. Certaines références sont aujourd'hui calculées avant le split : cette limite sera détaillée plus loin.", "Sources : notebook, cellules 12 à 28 ; docs/features.md")
    steps=[("FUSION","3 sources\n31 colonnes"),("NETTOYAGE","Types, pourcentages\net encodages"),("ENRICHISSEMENT","Rémunération relative\net satisfaction"),("MATRICE FINALE","43 variables\nnumériques")]
    for i,(title,body) in enumerate(steps):
        x=.55+i*3.12
        d.card(x,2.25,2.85,2.75,title,body,size=17)
        if i<3:d.arrow(x+2.88,3.60,x+3.08,3.60)
    d.banner("Contrat d'entrée : variables préparées, sans cible ni identifiant, avec les noms exacts de /features.")
    d.finish()

    d.new("Séparer l'évaluation du choix du seuil", "04 / Protocole", "La séparation stratifiée réserve 20 % des données pour le test : 294 salariés dont 47 départs. Le train compte 1 176 salariés dont 190 départs. Les modèles sont comparés par validation croisée à cinq folds ; la grille optimise le F1 et la calibration utilise sigmoid avec cinq folds. Le seuil est choisi sur les probabilités out-of-fold du train. Les transformations apprises avant le split limitent toutefois l'indépendance globale de l'évaluation actuelle.", "Sources : notebook, cellules 28 à 41 ; docs/model.md")
    d.box(.65,2.15,9.65,.66,TEAL,rounded=True)
    d.box(10.42,2.15,2.28,.66,NAVY,rounded=True)
    d.text("80 % : entraînement",.92,2.31,8.9,.3,17,WHITE,bold=True)
    d.text("20 % : test",10.6,2.31,1.95,.3,17,WHITE,bold=True)
    d.card(.65,3.2,5.85,2.35,"1 176 LIGNES","CV stratifiée à 5 folds\nOptimisation et seuil sur le train",size=17)
    d.card(6.72,3.2,5.98,2.35,"294 LIGNES","47 départs pour l'évaluation finale\nrandom_state = 42",size=17)
    d.banner("Une mesure sur les données d'entraînement ou des cibles fictives n'est pas une mesure de généralisation.")
    d.finish()

    d.new("Comparer plusieurs approches, sans surinterpréter le classement", "05 / Modèles comparés", "Cette comparaison utilise les F1 moyens en validation croisée enregistrés dans le notebook. LightGBM atteint ici le meilleur F1 moyen, tandis que la régression logistique est également compétitive. La forêt optimisée est le modèle exporté, mais il ne faut pas affirmer sa supériorité générale. Les écarts-types et les protocoles de choix des hyperparamètres doivent être pris en compte. Aucune comparaison finale sur le test de tous les concurrents avec des seuils optimisés identiquement n'est fournie.", "Sources : notebook, cellules 31, 34, 54 et 55 ; docs/model.md")
    models=[("LightGBM",.5226,.0426), ("Random Forest optimisé",.5048,.0714),("Régression logistique",.4968,.0280),("Random Forest + SMOTE",.4600,.0936),("Random Forest initial",.4189,.0967),("Dummy stratifié",.1644,.0300)]
    d.text("F1 moyen sur le train • validation croisée à 5 folds",.65,1.83,11.9,.35,16,MUTED)
    for i,(label,value,std) in enumerate(models):
        y=2.45+i*.52
        d.text(label,.65,y,3.7,.31,16,TEAL if i==1 else INK,bold=i==1)
        d.box(4.48,y+.015,6.0,.24,LINE)
        d.box(4.48,y+.015,6.0*value/.6,.24,TEAL if i==1 else "839DA6")
        d.text(f"{value:.4f} ± {std:.4f}".replace(".",","),10.65,y-.015,2.1,.36,14,INK,bold=i==1)
    d.text("0",4.47,5.75,.5,.24,11,MUTED)
    d.text("0,60",9.94,5.75,.55,.24,11,MUTED,align="right")
    d.banner("La forêt optimisée est le modèle livré ; les alternatives restent compétitives.")
    d.finish()

    d.new("Le modèle livré : forêt régularisée et probabilités calibrées", "06 / Modèle retenu", "La grille explore 16 combinaisons, évaluées sur cinq folds. Le modèle retenu utilise 300 arbres avec une profondeur maximale de cinq, un minimum de cinq observations pour une séparation et deux par feuille. La pondération des classes est balanced. La calibration sigmoid est réalisée avec CalibratedClassifierCV, cv=5. Le fichier complet conserve aussi le seuil, les noms de variables, la classe positive et les variables préparées des salariés pour l'import local.", "Sources : artefact local ; docs/model.md ; notebook, cellules 34 et 36")
    d.card(.55,2.05,3.93,3.55,"FORÊT","300 arbres\nProfondeur maximale : 5\nClasses pondérées\nrandom_state = 42",size=18)
    d.card(4.65,2.05,3.93,3.55,"CALIBRATION","Méthode sigmoid\nValidation interne : 5 folds\nProbabilité de la classe 1",size=18)
    d.card(8.75,2.05,3.99,3.55,"DÉCISION","Seuil optimisé : 0,2232\nChoisi sur le F1 OOF\nStocké avec le modèle",size=18)
    d.banner("L'artefact attend 43 variables déjà transformées ; la préparation brute n'est pas embarquée.")
    d.finish()

    d.new("Le seuil optimisé augmente le rappel sur le test", "07 / Performances", "Sur les 294 salariés du test, le rappel passe de 17,0 % à 61,7 % lorsque le seuil passe de 0,50 à environ 0,2232. Le F1 augmente de 0,229 à 0,500. La précision reste de 42,0 % : toutes les alertes ne sont pas de vrais départs. L'accuracy au seuil optimisé est de 80,27 % et la ROC-AUC de 0,7692. Une stratégie qui prédit toujours resté atteindrait déjà 84,0 % d'accuracy sur ce test : l'accuracy seule ne suffit pas. Les valeurs viennent du notebook, sans réévaluation pendant la fabrication du support.", "Source : notebook, cellule 41 ; docs/model_metrics.json • test : 294 lignes, 47 départs")
    for i,(key,label) in enumerate([("recall","RAPPEL"),("precision","PRÉCISION"),("f1","F1")]):
        x=.70+i*4.15
        d.text(label,x,1.92,3.8,.35,18,TEAL,bold=True)
        for j,(values,color,name) in enumerate([(native,"839DA6","Seuil 0,50"),(tuned,TEAL,"Seuil 0,2232")]):
            y=2.7+j*1.15
            d.text(name,x,y,3.55,.30,14,MUTED)
            d.box(x,y+.42,3.0,.31,LINE)
            d.box(x,y+.42,3.0*values[key],.31,color)
            d.text(f"{values[key]:.3f}".replace(".",","),x,y+.83,3.0,.4,21,color,bold=True)
    d.banner("Seuil optimisé : accuracy 0,8027  •  ROC-AUC 0,7692  •  Brier score 0,1226",y=5.85)
    d.finish()

    d.new("Chaque inférence suit un circuit contrôlé", "08 / Architecture", "Le client fournit du JSON, un CSV préparé ou un identifiant salarié. FastAPI et Pydantic valident le contrat. Le service enregistre les inputs en PostgreSQL, les relit dans une nouvelle session, puis appelle le modèle. Les outputs sont validés, committés et relus avant la réponse HTTP. Toutes les routes d'inférence passent par ce service. Ce choix ajoute des accès BDD mais garantit que les valeurs effectivement calculées sont celles stockées.", "Sources : api.py ; schemas.py ; inference.py ; docs/architecture.md")
    nodes=[(.65,2.0,"CLIENT","JSON / CSV / ID"),(4.8,2.0,"FASTAPI + PYDANTIC","Contrat et validation"),(8.95,2.0,"INPUTS EN BDD","Écriture puis relecture"),(8.95,4.3,"MODÈLE","predict / proba / score"),(4.8,4.3,"OUTPUTS EN BDD","Validation et persistance"),(.65,4.3,"RÉPONSE","Résultat relu depuis la BDD")]
    for x,y,title,body in nodes:d.card(x,y,3.72,1.75,title,body,size=15)
    d.arrow(4.39,2.82,4.70,2.82); d.arrow(8.54,2.82,8.85,2.82)
    d.arrow(10.81,3.80,10.81,4.20)
    d.arrow(8.85,5.12,8.54,5.12); d.arrow(4.70,5.12,4.39,5.12)
    d.finish()

    d.new("PostgreSQL structure les données et l'historique", "09 / Schéma relationnel", "La BDD contient huit tables. Les sources sont employees, surveys et evaluations, reliées par identifiant salarié. prepared_employees conserve les variables par salarié et version de modèle. model_versions conserve le SHA-256, les noms de variables et le seuil. api_exchanges relie chaque requête aux model_inputs et model_outputs. Une contrainte interdit de dupliquer une ligne d'entrée au sein d'un échange. La vue api_logs expose une lecture plus pratique. Le rôle applicatif ne peut pas modifier les sources ni supprimer les traces.", "Sources : database.py ; docs/database.md ; docs/schema.sql")
    d.card(.55,1.98,3.7,2.1,"SOURCES RH","employees\nsurveys • evaluations",size=16)
    d.card(4.8,1.98,3.7,2.1,"VARIABLES + VERSION","prepared_employees\nmodel_versions",size=16)
    d.arrow(4.31,3.04,4.72,3.04)
    d.card(9.03,1.98,3.7,2.1,"ÉCHANGE API","api_exchanges\nID, route, statut, durée",size=16)
    d.arrow(8.55,3.04,8.95,3.04)
    d.card(4.8,4.48,3.7,1.77,"ENTRÉES","model_inputs\nfeatures, cible, ordre",size=15)
    d.card(9.03,4.48,3.7,1.77,"SORTIES","model_outputs\nréponse et valeurs modèle",size=15)
    d.arrow(10.86,4.12,10.86,4.39)
    d.arrow(9.7,4.14,7.72,4.40)
    d.text("8 tables\n+ vue api_logs",.70,4.78,3.4,1.20,27,TEAL,bold=True)
    d.finish()

    d.new("Une requête peut être expliquée et retrouvée", "10 / Traçabilité", "Chaque requête reçoit un identifiant unique X-Request-ID. Le journal conserve la route, le contexte utile, les dates, la durée, le fichier, le nombre de lignes, la version du modèle et les erreurs. Les étapes persistées permettent de voir où un traitement s'est arrêté. Un échec du modèle conserve les inputs. Si la BDD n'est pas accessible dès le départ, l'API refuse le calcul et l'erreur ne peut être écrite que dans les logs du processus. Un arrêt brutal peut laisser pending ; il n'y a pas de garantie exactement une fois.", "Sources : audit.py ; api.py ; docs/api_logs.md")
    d.text("X-Request-ID relie le journal, les entrées, la version et les sorties.",.65,1.89,12,.55,19,INK,bold=True)
    stages=["received","inputs_saved","model_running","outputs_saved","completed"]
    for i,stage in enumerate(stages):
        x=.65+i*2.45
        d.box(x,2.85,2.23,.75,TEAL if i==4 else WHITE,rounded=True)
        d.text(stage,x+.12,3.08,1.99,.31,13,WHITE if i==4 else INK,align="center",bold=True)
        if i<4:d.arrow(x+2.25,3.22,x+2.40,3.22)
    d.card(.65,4.1,5.92,1.98,"SUCCÈS","Message, durée, lignes, fichier et résumé du résultat.",size=17)
    d.card(6.79,4.1,5.92,1.98,"ÉCHEC","Cause détaillée et dernière étape atteinte, consultables dans api_logs.",size=17)
    d.finish()

    d.new("Trois routes CSV au cœur d'une API de sept endpoints", "11 / API", "Présenter les trois routes demandées pour le fichier : predict/file renvoie les classes natives ; predict-proba/file renvoie les probabilités dans l'ordre de classes ; score/file exige les vraies réponses et renvoie l'accuracy native et les métriques au seuil optimisé. Les routes JSON et par ID utilisent le seuil optimisé. Les deux endpoints techniques sont health et features. L'ordre des lignes est conservé. Les CSV doivent être UTF-8 et leur contenu est limité à dix Mio. Swagger permet de sélectionner un fichier et de tester une route.", "Sources : api.py ; docs/api.md ; docs/openapi.json")
    d.table(["Méthode / route","Usage"],[
        ("POST /predict/file","Classes natives par ligne"),
        ("POST /predict-proba/file","Probabilités par classe"),
        ("POST /score/file","Score avec les vraies réponses"),
        ("POST /predict","JSON au seuil optimisé"),
        ("GET /employees/{id}/score","Variables du salarié lues en BDD"),
        ("GET /health • GET /features","Santé et noms des variables"),
    ],.65,1.95,[5.65,6.4],row_h=.56,font=15)
    d.banner("predict/file applique la décision native ; predict et la route par ID appliquent le seuil optimisé.",y=6.14)
    d.finish()

    d.new("Pydantic valide les entrées et les sorties", "12 / Contrats de données", "La validation est partagée entre le JSON et chaque ligne CSV. Elle contrôle les noms de variables selon le modèle, la présence de valeurs numériques finies et les cibles binaires. Les booléens et valeurs nulles ne sont pas acceptés comme nombres. Les sorties sont contrôlées avant insertion : classe entière 0 ou 1, probabilités bornées dont la somme vaut un à la tolérance près, dimensions cohérentes et métriques bornées. Pydantic assure le respect du contrat, pas la performance statistique ni toutes les règles métier.", "Sources : schemas.py ; docs/api_validation.md")
    d.card(.65,2.0,5.9,3.65,"AVANT LE CALCUL","Variables exactes et dictionnaire non vide\n\nNombres finis, sans booléens ni null\n\nCible 0/1 ou Non/Oui ; ID positif",size=18)
    d.card(6.78,2.0,5.9,3.65,"AVANT L'ENREGISTREMENT","Classes entières 0 ou 1\n\nProbabilités dans [0 ; 1], somme = 1\n\nUne sortie par ligne ; métriques bornées",size=18)
    d.banner("422 : entrée invalide  •  500 : calcul ou sortie invalide  •  503 : BDD ou modèle indisponible")
    d.finish()

    d.new("Tester le service sans confondre tests et performance ML", "13 / Qualité logicielle", "La suite vérifiée contient 87 tests et la couverture des modules suivis est arrondie à 89 %. Chaque endpoint dispose de cas valides et invalides. Les modèles simulés permettent de fixer précisément les réponses attendues. Les tests de BDD vérifient la relecture des inputs committés avant calcul, les sorties persistées et le comportement en cas d'échec. La CI utilise PostgreSQL avec un schéma isolé par test. Ces tests logiciels ne constituent pas une mesure de généralisation du modèle.", "Sources : tests/ ; rapports de vérification du projet ; docs/api_validation.md")
    d.stat(.70,1.95,3.7,"87","tests automatisés")
    d.stat(4.85,1.95,3.7,"89 %","de couverture des modules suivis")
    d.stat(9.00,1.95,3.7,"7","endpoints testés")
    d.card(.65,4.0,5.92,1.88,"API","Réponses attendues, erreurs, paramètres et validation des fichiers.",size=17)
    d.card(6.79,4.0,5.92,1.88,"BASE DE DONNÉES","Import idempotent, inputs avant calcul et refus des sorties non persistées.",size=17)
    d.finish()

    d.new("Le push déclenche une publication contrôlée", "14 / CI/CD et Hugging Face", "À chaque push ou pull request, la pipeline vérifie les tests avec PostgreSQL, la documentation et le package. L'export préparé ne conserve que model, threshold, feature_names et positive_class. Les lignes individuelles ne sont pas publiées avec cet export. Sur main, après validation d'environnement, le job transmet les quatre fichiers à SamGabb/hrnova avec le secret HF_TOKEN. Cela ne réentraîne pas le modèle et ne crée pas une API hébergée. Les protections de branche doivent être configurées dans GitHub pour bloquer les fusions.", "Sources : .github/workflows/ci.yml ; docs/huggingface.md ; docs/deployment.md")
    labels=[("PUSH / PR","Code et artefact"),("VÉRIFICATIONS","Tests, docs, package"),("EXPORT VALIDÉ","Sans lignes salariées"),("MAIN → HUB","SamGabb/hrnova")]
    for i,(title,body) in enumerate(labels):
        x=.65+i*3.08
        d.card(x,2.20,2.82,2.35,title,body,size=16)
        if i<3:d.arrow(x+2.85,3.35,x+3.03,3.35)
    d.text("Production : validations et éventuelles approbations GitHub.",.72,5.05,11.9,.45,19,TEAL,bold=True)
    d.banner("Publication du modèle ≠ réentraînement automatique ≠ déploiement d'une API publique.",y=6.02)
    d.finish()

    d.new("Une documentation pour utiliser et maintenir le projet", "15 / Transmission", "La documentation comporte douze pages MkDocs avec navigation et recherche. Elle distingue le parcours utilisateur, les contrats API et les informations développeur. Swagger et ReDoc décrivent les endpoints ; le contrat OpenAPI peut être exporté sans connexion BDD. Les guides détaillent les performances historiques et leurs limites, le schéma, les logs, l'installation, le déploiement et le protocole de maintenance. La CI construit le site en mode strict et conserve l'HTML comme artefact ; elle ne publie pas automatiquement un site public.", "Sources : mkdocs.yml ; docs/index.md ; README.md")
    d.card(.55,2.05,3.93,3.35,"UTILISATEUR","Swagger et ReDoc\nExemples JSON / CSV\nContrat des 43 variables\nCodes d'erreur",size=18)
    d.card(4.65,2.05,3.93,3.35,"DÉVELOPPEUR","Architecture et schéma\nTests et logs\nVersions et déploiement\nSources des métriques",size=18)
    d.card(8.75,2.05,3.99,3.35,"MAINTENEUR","Mises à jour du modèle\nSauvegardes et restauration\nRetour arrière\nDiagnostic des incidents",size=18)
    d.banner("Swagger : localhost:8000/docs  •  Documentation MkDocs : localhost:8001")
    d.finish()

    d.new("Les limites définissent les prochaines priorités", "16 / Regard critique", "La première priorité est d'apprendre les transformations uniquement sur le train et dans chaque fold, puis de sérialiser l'ensemble de la préparation. Aujourd'hui, les médianes et l'encodeur nominal sont calculés avant le split ; cela peut influencer les résultats. Il faut ensuite confirmer les performances sur des données temporelles ou externes, car le test actuel ne comporte que 47 départs. Une analyse des erreurs par sous-population est nécessaire. Enfin, le service est local et sans authentification HTTP : son exposition distante demanderait un travail spécifique.", "Sources : docs/model.md ; docs/maintenance.md ; docs/deployment.md")
    items=[("01","Préparation sans fuite","Ajuster médianes et encodeurs dans les folds ; exporter une pipeline complète."),("02","Validation indépendante","Évaluer sur des données nouvelles ; documenter incertitude et calibration."),("03","Usage responsable","Étudier les erreurs par groupe et définir les coûts des fausses alertes."),("04","Exploitation distante","Prévoir authentification, accès BDD, supervision et déploiement sécurisé.")]
    for i,(number,title,body) in enumerate(items):
        y=1.90+i*1.10
        d.text(number,.68,y,.67,.55,26,TEAL,bold=True)
        d.text(title,1.65,y,10.8,.36,20,INK,bold=True)
        d.text(body,1.65,y+.46,10.8,.49,17,MUTED)
    d.finish()

    d.new("Un modèle exploitable, un résultat vérifiable", "17 / Conclusion", "Conclure sur les éléments réellement livrés : modèle calibré, API validée, données et échanges structurés en BDD, tests, documentation et publication contrôlée du modèle. Le rappel augmente avec le seuil optimisé, mais les performances demeurent limitées et doivent être confirmées après correction de la préparation. La suite logique est une pipeline de prétraitement réutilisable, une validation indépendante et une exploitation surveillée. Inviter aux questions ou proposer la démonstration en annexe.", "Sources : documentation du projet et résultats vérifiés", dark=True)
    d.text("ESTIMER",.70,2.1,3.7,.55,29,LIME,bold=True)
    d.text("Une probabilité calibrée\net un seuil explicite.",.70,3.0,3.7,1.4,22,WHITE)
    d.text("TRACER",4.9,2.1,3.7,.55,29,LIME,bold=True)
    d.text("Des inputs, des outputs\net une version retrouvables.",4.9,3.0,3.7,1.4,22,WHITE)
    d.text("MAINTENIR",9.05,2.1,3.7,.55,29,LIME,bold=True)
    d.text("Des tests, des guides\net un protocole de mise à jour.",9.05,3.0,3.7,1.4,22,WHITE)
    d.text("Questions / démonstration",.70,5.65,11.7,.65,25,WHITE,bold=True)
    d.finish()

    d.new("Démonstration : du fichier au journal de requête", "Annexe A / Démo locale", "Avant la soutenance, démarrer PostgreSQL et l'API. Ouvrir Swagger pour montrer les sept endpoints. Envoyer le CSV synthétique aux routes predict-proba/file puis score/file. Expliquer que les cibles de l'exemple sont fictives et que le score obtenu n'est pas une mesure de performance. Relever le X-Request-ID et retrouver le journal, les inputs et les outputs dans PostgreSQL. Terminer avec une requête invalide pour montrer le code 422 et la cause dans api_logs. Ne jamais afficher .env ou les tokens pendant la démonstration.", "Sources : examples/ ; docs/api.md ; docs/api_logs.md")
    d.card(.65,1.97,5.9,3.9,"1. APPELER L'API","Ouvrir localhost:8000/docs\n\nPOST /predict-proba/file\nFichier : examples/employees.csv\n\nPOST /score/file\nFichier : employees_with_target.csv",size=16)
    d.card(6.78,1.97,5.9,3.9,"2. RETROUVER LA TRACE","Relever X-Request-ID\n\nOuvrir public.api_logs\nFiltrer avec requete_id\n\nContrôler route, durée, étape,\nmessage et résumé du résultat",size=17)
    d.banner("Les exemples sont synthétiques : les scores de démonstration ne sont pas des performances ML.",y=6.10)
    d.finish()

    d.new("Sources et éléments de preuve", "Annexe B / Références", "Les chiffres du modèle proviennent des sorties enregistrées dans le notebook et du relevé model_metrics.json, associé au hash de l'artefact. Le rapport de tests prouve la conformité logicielle, pas la performance statistique. Le contrat OpenAPI et la documentation sont construits depuis les sources. Les liens GitHub et Hugging Face permettent de retrouver le projet et la version publiée ; leur accessibilité dépend des droits du visiteur.", "Références du projet • état examiné le 27 octobre 2026")
    refs=[("DONNÉES ET MODÈLE","dataScientistTechNova.ipynb • cellules 28 à 41\ndocs/model.md • docs/model_metrics.json"),("API, BDD ET VALIDATION","docs/api.md • docs/api_validation.md\ndocs/database.md • docs/api_logs.md"),("EXPLOITATION ET SOURCES","docs/maintenance.md • docs/deployment.md\ngithub.com/samuelGabbai/P5_Futurisys\nhuggingface.co/SamGabb/hrnova")]
    for i,(title,body) in enumerate(refs):
        y=1.9+i*1.52
        d.box(.65,y,12.05,1.35,WHITE,rounded=True)
        d.text(title,.89,y+.13,11.55,.35,15,TEAL,bold=True)
        d.text(body,.89,y+.52,11.55,.75,14)
    d.finish()
    assert len(d.images)==20
    d.save(out)


if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir",type=Path,default=ROOT/"presentations")
    build(parser.parse_args().output_dir)
