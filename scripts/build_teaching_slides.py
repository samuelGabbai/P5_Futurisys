"""Créer les slides pédagogiques et enrichir une copie de la soutenance.

uv run --no-project --with python-pptx==1.0.2 python scripts/build_teaching_slides.py
"""

from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import re

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from build_presentation import Deck, ROOT, NAVY, TEAL, WHITE, INK, PALE

OUT = ROOT / "presentations"
STEM = "HRNova_Git_Pydantic_PostgreSQL_2026-10-29"
GIT = "https://git-scm.com/book/en/v2/Getting-Started-What-is-Git%3F"
PYDANTIC = "https://docs.pydantic.dev/latest/concepts/validators/"
PG = "https://www.postgresql.org/docs/18/tutorial-transactions.html"
ACTIONS = "https://docs.github.com/en/actions/get-started/understand-github-actions"


def code(d, text, x, y, w, h, size=17):
    d.box(x, y, w, h, NAVY, rounded=True)
    d.text(text, x + .22, y + .22, w - .44, h - .44, size, WHITE)


def steps(d, labels, y=2.05):
    width = (12.2 - .28 * (len(labels) - 1)) / len(labels)
    for i, (title, body) in enumerate(labels):
        x = .55 + i * (width + .28)
        d.card(x, y, width, 2.55, title, body, size=17)
        if i + 1 < len(labels):
            d.arrow(x + width + .03, y + 1.25, x + width + .25, y + 1.25)


def build():
    d = Deck(total=12)
    d.prs.core_properties.title = "HRNova — Comprendre Git, Pydantic et PostgreSQL"
    d.prs.core_properties.created = datetime(2026, 10, 29)

    d.new("Git garde l’histoire ; GitHub la partage", "Git · 1 / 4",
          "Git est le logiciel de gestion de versions. Il fonctionne en local : chaque commit conserve un état du projet et ses liens avec les commits précédents. GitHub héberge une copie distante du dépôt et apporte les pull requests. GitHub Actions exécute des workflows sur des machines appelées runners. Un commit local ne lance pas à lui seul les tests sur GitHub : il faut notamment un push et un workflow configuré pour cet événement.", GIT)
    steps(d, [("Git", "Sur ton ordinateur\nHistorique des fichiers\nCommits et branches"),
              ("GitHub", "Dépôt distant\nPartage du code\nPull requests et revues"),
              ("GitHub Actions", "Automatisation\nTests à chaque push\nPublication du modèle")])
    d.banner("Dans HRNova : je modifie le code → je pousse → les contrôles se lancent.", y=5.55)
    d.finish()

    d.new("Du fichier modifié au dépôt distant", "Git · 2 / 4",
          "Le répertoire de travail contient les fichiers que l’on modifie. git add sélectionne les changements à enregistrer et les place dans l’index. git commit crée un instantané local avec un message, un auteur et un parent. git push transmet les commits à la branche distante. git status permet de vérifier chaque étape. Une fois un fichier ajouté, une nouvelle modification doit être ajoutée à nouveau pour entrer dans le prochain commit. Ne jamais versionner de jetons ni le fichier .env.", GIT)
    steps(d, [("1. Modifier", "Répertoire de travail\nEx. : schemas.py"),
              ("2. Sélectionner", "Index / staging\nChoisir les changements"),
              ("3. Enregistrer", "Historique local\nCréer un commit"),
              ("4. Partager", "Dépôt distant\nEnvoyer les commits")])
    code(d, 'git status\ngit add schemas.py\ngit commit -m "Valider les entrees"\ngit push', .55, 4.95, 7.35, 1.57, 16)
    d.text(".gitignore écarte .env des nouveaux suivis.\nIl ne supprime pas un secret déjà commité.", 8.2, 5.1, 4.4, 1.1, 17)
    d.finish()

    d.new("Branches : proposer, vérifier, puis fusionner", "Git · 3 / 4",
          "Une branche est une référence vers un commit, pas une copie indépendante du dossier. Créer une branche permet de développer une fonctionnalité sans faire évoluer main immédiatement. Après le push, la pull request présente les différences, les contrôles et les commentaires. Les conflits surviennent lorsque Git ne peut pas combiner automatiquement des modifications ; il faut choisir le contenu correct et enregistrer la résolution. Les vérifications obligatoires et la protection de main se configurent dans GitHub : le YAML seul ne bloque pas les fusions.", GIT)
    code(d, "git switch -c feature/validation\n# modifier puis commiter\ngit push -u origin feature/validation", .55, 1.95, 12.2, 1.55, 19)
    steps(d, [("Branche", "Développer une évolution\nHistorique isolé de main"),
              ("Pull request", "Comparer les changements\nTests et revue du code"),
              ("Fusion", "Intégrer dans main\nRésoudre les conflits")], y=3.75)
    d.finish()

    d.new("GitHub Actions : que se passe-t-il au push ?", "Git · 4 / 4",
          "Un workflow est un fichier YAML sous .github/workflows. Il définit les événements, les jobs et leurs étapes. Chaque job s’exécute sur un runner ; les étapes d’un même job s’enchaînent. Les jobs peuvent fonctionner en parallèle, sauf dépendance déclarée avec needs. Dans HRNova, la CI vérifie le projet et l’intégration PostgreSQL. La publication Hugging Face envoie les artefacts préparés vers SamGabb/hrnova ; elle ne réentraîne pas automatiquement le modèle. Le jeton est fourni par les secrets GitHub au job de publication, jamais écrit dans le dépôt.", ACTIONS)
    steps(d, [("Événement", "push ou pull_request\nDéclenche le workflow"),
              ("Runner", "Machine de travail\nInstalle Python et outils"),
              ("Jobs / étapes", "Exécute les contrôles\nSignale succès ou échec")])
    code(d, ".github/workflows/ci.yml\nContrôles → tests → artefacts → publication autorisée", .55, 4.9, 12.2, 1.08, 18)
    d.banner("Publier un modèle existant et réentraîner un modèle sont deux opérations distinctes.", y=6.15)
    d.finish()

    d.new("Pydantic transforme une entrée en contrat", "Pydantic · 1 / 4",
          "Pydantic décrit la structure et les types attendus. Il vérifie la donnée et peut la convertir selon les règles choisies. Les exemples de cette diapositive montrent un seul champ numérique : une vraie requête HRNova doit également fournir toutes les variables attendues. Le projet accepte une chaîne numérique convertible et rejette les booléens, les valeurs absentes et les nombres non finis. Certaines de ces règles sont des validateurs spécifiques à HRNova ; elles ne sont pas toutes les valeurs par défaut de Pydantic. La validation technique ne garantit pas la qualité statistique de la prédiction.", PYDANTIC)
    d.table(["Valeur reçue pour un champ", "Résultat dans HRNova"],
            [("35", "Acceptée → 35.0"), ('"35"', "Convertie → 35.0"),
             ('"trente"', "Rejetée → HTTP 422"), ("null ou true", "Rejeté → HTTP 422")],
            .55, 1.95, [5.7, 6.5], row_h=.70, font=19)
    d.banner("Validation = structure + types + règles. Une entrée valide peut encore être peu réaliste.", y=6.0)
    d.finish()

    d.new("Un schéma Python décrit les données attendues", "Pydantic · 2 / 4",
          "Cet extrait est volontairement simplifié : la version complète est dans schemas.py. NumericValue associe un nombre fini à un validateur exécuté avant la conversion. Le validateur numeric_value rejette notamment les booléens et les valeurs nulles. Field(min_length=1) impose au moins une variable. ConfigDict(extra='forbid') rejette les champs supplémentaires à la racine du modèle. Un autre validateur du projet compare les clés de features à la liste exacte des variables attendues : il détecte les variables manquantes et inconnues. Ce dernier contrôle n’est pas assuré par extra='forbid' dans un dictionnaire.", "Code du projet : schemas.py · " + PYDANTIC)
    code(d, 'NumericValue = Annotated[\n    FiniteFloat, BeforeValidator(numeric_value)\n]\n\nclass PredictionRequest(BaseModel):\n    model_config = ConfigDict(extra="forbid")\n    features: dict[str, NumericValue] = Field(min_length=1)', .55, 1.95, 8.0, 3.75, 17)
    d.text("1. Contrôler les valeurs\n\n2. Convertir les nombres\n\n3. Vérifier les noms attendus", 8.9, 2.2, 3.7, 2.95, 18)
    d.banner("Extrait simplifié : les clés de features sont aussi vérifiées par un validateur dédié.", y=6.0)
    d.finish()

    d.new("Un fichier CSV passe les mêmes contrôles", "Pydantic · 3 / 4",
          "Le fichier envoyé à l’API est décodé puis transformé en lignes. Chaque ligne est soumise au schéma de prédiction ; le lot doit contenir au moins une ligne. Le projet vérifie les variables attendues, convertit les valeurs numériques et prépare les colonnes dans l’ordre du modèle. Une erreur est renvoyée avec un contexte exploitable pour retrouver la ligne et le champ. Le modèle n’est pas appelé si l’entrée est invalide. Pour score, les cibles doivent aussi être valides et alignées avec les lignes : les formes acceptées sont Non/Oui ou 0/1 selon les règles du projet.", "Code du projet : schemas.py, inference.py")
    steps(d, [("Fichier CSV", "Lire le contenu\nRepérer les colonnes"),
              ("Lignes typées", "Valider chaque ligne\nVérifier chaque variable"),
              ("Entrées valides", "Ordonner les colonnes\nEnregistrer puis calculer")])
    d.box(.55, 4.95, 12.2, 1.1, PALE, rounded=True)
    d.text('Exemple : une cellule numérique contient "trente".', .8, 5.1, 11.7, .35, 19, TEAL, bold=True)
    d.text("Réponse 422 avec le contexte de l’erreur ; aucune prédiction produite.", .8, 5.56, 11.7, .35, 17)
    d.finish()

    d.new("Valider aussi ce que renvoie le modèle", "Pydantic · 4 / 4",
          "La validation continue après le calcul. Le projet impose des classes binaires 0 et 1 pour les prédictions, des probabilités finies dans l’intervalle [0,1] et une somme par ligne égale à 1 à une tolérance de 10 puissance moins 6. Les réponses par fichier doivent avoir le bon nombre de lignes. Les métriques sont également contrôlées. Une sortie invalide provoque une erreur 500 et une trace d’échec ; les entrées déjà enregistrées restent disponibles, mais aucune sortie invalide n’est sauvegardée comme résultat. Pydantic contrôle le contrat de réponse, pas la pertinence métier ni la performance sur de nouvelles données.", "Code du projet : schemas.py, inference.py")
    for x, title, body in [(.55, "predict", "Classe autorisée : 0 ou 1\nUne prédiction par ligne"),
                           (4.72, "predict_proba", "Probabilités entre 0 et 1\nSomme par ligne ≈ 1"),
                           (8.89, "score", "Métriques finies\nValeurs dans les bornes")]:
        d.card(x, 2.05, 3.85, 2.8, title, body, size=18)
    d.banner("Une réponse hors contrat est rejetée et tracée ; elle ne devient pas un résultat valide.", y=5.65)
    d.finish()

    d.new("PostgreSQL : où vivent les données ?", "PostgreSQL · 1 / 4",
          "PostgreSQL est le serveur qui stocke les données et exécute les requêtes SQL. Dans la configuration locale du projet, le serveur est joint sur 127.0.0.1 au port 5433 et la base s’appelle hrnova. Une base contient des schémas, dont public ; les schémas contiennent des tables. Une table possède des colonnes typées et des lignes. DBeaver est un client pour se connecter au serveur et explorer ces objets. Installer DBeaver ne démarre pas PostgreSQL et ne crée pas la base. Les paramètres de connexion restent configurables par l’environnement ; aucun mot de passe n’est inclus ici.", "Configuration et modèles de données du projet")
    steps(d, [("Serveur", "PostgreSQL local\n127.0.0.1:5433"),
              ("Base / schéma", "Base : hrnova\nSchéma : public"),
              ("Tables", "employees\nmodel_inputs / outputs")])
    d.box(.55, 5.0, 12.2, 1.1, PALE, rounded=True)
    d.text("DBeaver → connexion → PostgreSQL → tables", .8, 5.18, 11.7, .4, 23, TEAL, bold=True)
    d.text("DBeaver affiche les données ; le serveur PostgreSQL doit être démarré.", .8, 5.67, 11.7, .3, 16)
    d.finish()

    d.new("Les clés relient les tables entre elles", "PostgreSQL · 2 / 4",
          "Une clé primaire identifie une ligne de manière unique. Une clé étrangère pointe vers une ligne d’une autre table et empêche de créer un lien vers un identifiant inexistant. L’exemple emploie l’identifiant fictif 42. Dans HRNova, employees.id_employee est relié aux tables d’enquêtes et d’évaluations par employee_id. Les échanges API possèdent également leur identifiant ; les entrées et les sorties y sont rattachées. La paire exchange_id et row_index est unique pour les entrées d’un échange. SQLAlchemy traduit les objets Python en SQL, puis PostgreSQL applique les contraintes. JSONB permet de stocker les variables et résultats structurés.", "Modèles SQLAlchemy du projet")
    d.card(.55, 2.0, 3.7, 2.8, "employees", "id_employee = 42\nClé primaire : identité", size=18)
    d.card(4.8, 2.0, 3.7, 2.8, "surveys", "employee_id = 42\nClé étrangère : lien", size=18)
    d.card(9.05, 2.0, 3.7, 2.8, "evaluations", "employee_id = 42\nClé étrangère : lien", size=18)
    d.arrow(4.3, 3.35, 4.75, 3.35)
    d.arrow(4.4, 4.95, 10.8, 4.95)
    d.text("Exemple fictif : les tables partagent un identifiant, sans recopier toute la fiche.", .55, 5.2, 12.2, .6, 18)
    d.banner("SQLAlchemy : objets Python → requêtes SQL → contraintes vérifiées par PostgreSQL.", y=6.05)
    d.finish()

    d.new("Une transaction protège un groupe d’écritures", "PostgreSQL · 3 / 4",
          "Une transaction regroupe des instructions SQL : commit les rend durables, rollback les annule en cas d’échec. Les garanties ACID couvrent l’atomicité, la cohérence des contraintes, l’isolation et la durabilité. Elles s’appliquent à chaque transaction, pas automatiquement à l’ensemble d’une requête HTTP. Dans HRNova, les entrées sont validées et commitées avant le calcul, puis relues depuis la base. Le modèle produit une sortie qui est validée, enregistrée puis relue. Cela conserve les entrées en cas d’échec ultérieur. Si PostgreSQL est indisponible, l’inférence ne doit pas contourner la base ; si la sortie ne peut pas être enregistrée, elle n’est pas livrée comme un succès.", PG)
    steps(d, [("Transaction A", "Sauver les entrées\nCOMMIT avant le modèle"),
              ("Calcul", "Relire depuis la base\nExécuter le modèle"),
              ("Transaction B", "Valider et sauver la sortie\nCOMMIT puis relecture")])
    code(d, "with session() as db, db.begin():\n    db.add(entree)\n# succès : commit ; exception : rollback", .55, 4.95, 8, 1.35, 17)
    d.text("Les entrées déjà commitées restent traçables si le calcul échoue.", 8.9, 5.12, 3.7, 1.12, 18)
    d.finish()

    d.new("Tracer un appel : retrouver son parcours en SQL", "PostgreSQL · 4 / 4",
          "api_exchanges décrit la requête et son déroulement. model_inputs conserve les entrées validées, model_outputs les résultats valides, et model_versions identifie le modèle. L’identifiant de requête permet de corréler ces tables et les logs. La vue api_logs présente une lecture simplifiée des échanges ; une vue est une requête sauvegardée, pas une duplication des traces. Les requêtes rejetées gardent une trace d’erreur, même sans entrée modèle valide. Le rôle applicatif doit disposer des seuls droits nécessaires. Un échec total de la base empêche évidemment d’écrire cette trace en base : les logs applicatifs complètent alors l’investigation.", "Schéma PostgreSQL du projet · vue public.api_logs")
    code(d, "SELECT date_requete, route, code_http, message\nFROM public.api_logs\nORDER BY date_requete DESC\nLIMIT 5;", .55, 1.95, 12.2, 1.9, 20)
    d.card(.55, 4.15, 5.95, 2.35, "Une requête, un identifiant", "Relier l’appel API, les entrées, les sorties et les étapes d’exécution.", size=17)
    d.card(6.8, 4.15, 5.95, 2.35, "Un modèle identifiable", "Retrouver la version du modèle utilisée et comprendre les erreurs.", size=17)
    d.finish()
    return d


def save_teaching(d):
    OUT.mkdir(exist_ok=True)
    d.prs.save(OUT / f"{STEM}.pptx")
    overview = Image.new("RGB", (1600, 675), "white")
    previews = OUT / "apercus_pedagogiques"
    previews.mkdir(exist_ok=True)
    notes = ["# Git, Pydantic et PostgreSQL — Notes pour l’oral", "", "12 diapositives. Les notes sont aussi intégrées au PowerPoint.", ""]
    for i, (img, (title, note, source)) in enumerate(zip(d.images, d.notes), 1):
        overview.paste(img.resize((400, 225)), (((i - 1) % 4) * 400, ((i - 1) // 4) * 225))
        img.save(previews / f"slide_{i:02d}.png")
        notes.extend([f"## {i:02d} — {title}", "", note, "", f"Source : {source}", ""])
    overview.save(OUT / f"{STEM}_Apercu.png")
    (OUT / f"{STEM}_Notes.md").write_text("\n".join(notes), encoding="utf-8")


def enrich(d):
    original = OUT / "HRNova_Soutenance_2026-10-27.pptx"
    if not original.exists():
        return
    digest = sha256(original.read_bytes()).hexdigest()
    prs = Presentation(original)
    assert len(prs.slides) == 20, "La soutenance source doit contenir 20 diapositives."
    baseline = [tuple(s.text for s in slide.shapes if s.has_text_frame and not re.fullmatch(r"\d+\s*/\s*\d+", s.text.strip())) for slide in prs.slides]
    old_ids = list(prs.slides._sldIdLst)
    for src in d.prs.slides:
        assert not src._element.xpath(".//@r:embed"), "Les slides ajoutées doivent être autonomes."
        dest = prs.slides.add_slide(prs.slide_layouts[6])
        for shape in list(dest.shapes):
            dest.shapes._spTree.remove(shape._element)
        for shape in src.shapes:
            dest.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
        dest.background.fill.solid()
        dest.background.fill.fore_color.rgb = src.background.fill.fore_color.rgb
        # Le masque de notes du fichier édité peut ne plus avoir de zone de texte.
        if dest.notes_slide.notes_text_frame is None:
            for shape in src.notes_slide.shapes:
                if shape.is_placeholder and shape.has_text_frame and shape.text_frame == src.notes_slide.notes_text_frame:
                    dest.notes_slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
            if dest.notes_slide.notes_text_frame is None:
                # Copier le placeholder BODY en s'appuyant sur son type XML.
                for shape in src.notes_slide.shapes:
                    if shape.is_placeholder and shape._element.xpath('.//p:ph[@type="body"]'):
                        dest.notes_slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
        dest.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    new_ids = list(prs.slides._sldIdLst)[20:]
    order = []
    for i, element in enumerate(old_ids, 1):
        order.append(element)
        order.extend({10: new_ids[8:12], 13: new_ids[4:8], 15: new_ids[0:4]}.get(i, []))
    for element in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(element)
    for element in order:
        prs.slides._sldIdLst.append(element)
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.left > Inches(10) and re.fullmatch(r"\d+\s*/\s*\d+", shape.text.strip()):
                runs = [run for p in shape.text_frame.paragraphs for run in p.runs]
                if runs:
                    runs[0].text = f"{i:02d} / 32"
                    for run in runs[1:]:
                        run.text = ""
    prs.core_properties.title = "HRNova — Soutenance enrichie : Git, Pydantic et PostgreSQL"
    prs.core_properties.modified = datetime(2026, 10, 29)
    target = OUT / "HRNova_Soutenance_Enrichie_2026-10-29.pptx"
    prs.save(target)
    check = Presentation(target)
    assert len(check.slides) == 32
    original_positions = [i for i, element in enumerate(order) if element in old_ids]
    for original_texts, position in zip(baseline, original_positions):
        actual = tuple(s.text for s in check.slides[position].shapes if s.has_text_frame and not re.fullmatch(r"\d+\s*/\s*\d+", s.text.strip()))
        assert original_texts == actual, "Un texte de la soutenance source a été modifié."
    assert sha256(original.read_bytes()).hexdigest() == digest, "Le fichier source doit rester inchangé."
    print(f"Soutenance enrichie : {target.name} (32 slides ; original conservé)")


def main():
    d = build()
    assert len(d.prs.slides) == 12
    for slide in d.prs.slides:
        assert slide.notes_slide.notes_text_frame.text.strip()
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= d.prs.slide_width + 10
            assert shape.top + shape.height <= d.prs.slide_height + 10
    save_teaching(d)
    enrich(d)
    print(f"Support pédagogique : {STEM}.pptx (12 slides, notes et aperçus)")


if __name__ == "__main__":
    main()
