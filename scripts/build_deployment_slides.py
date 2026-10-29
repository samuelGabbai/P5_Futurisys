"""Support pédagogique complémentaire : déployer par environnement.

uv run --no-project --with python-pptx==1.0.2 python scripts/build_deployment_slides.py
"""

from copy import deepcopy
from hashlib import sha256
import re

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from build_presentation import Deck, ROOT
from build_teaching_slides import code, steps

OUT = ROOT / "presentations"
STEM = "HRNova_Deploiement_Environnements_2026-10-29"


def build():
    d = Deck(total=4)
    d.prs.core_properties.title = "HRNova — Déploiement par environnement"
    d.new("Un environnement = une instance et sa configuration", "Déploiement · 1 / 4",
          "Un environnement associe le code de l’API, sa configuration, sa base PostgreSQL et une version du modèle. Development sert à développer, staging à répéter la mise en service et production à servir les utilisateurs. Chaque cible doit disposer de ses accès et de sa base, pour éviter que les tests écrivent dans les données de production. Le code peut être identique ; les paramètres changent. Dans HRNova, l’API actuellement disponible reste l’instance locale. Les trois environnements GitHub existants organisent les validations sans créer trois serveurs.",
          "Sources : docs/deployment.md · .github/workflows/ci.yml")
    steps(d, [("Development", "Développer et déboguer\nDonnées de test\nRetours rapides"),
              ("Staging", "Préproduction\nTester le candidat\nConfiguration réaliste"),
              ("Production", "Servir les utilisateurs\nDonnées et accès dédiés\nVersion validée")])
    d.banner("Séparer les bases, les secrets et les instances ; identifier les versions du code et du modèle.", y=5.7)
    d.finish()

    d.new("Du push à la mise en production", "Déploiement · 2 / 4",
          "Le parcours cible commence par une branche de travail et des contrôles CI. Une pull request permet de revoir les changements. Le candidat est testé en préproduction avant une validation de mise en production. L’objectif est de promouvoir les mêmes artefacts validés, identifiés par un commit Git et un hash de modèle. Si la fusion change le code, le nouveau candidat doit être vérifié. Après déploiement, contrôler la santé de l’API, une prédiction et sa trace en base. Ce protocole décrit une cible d’exploitation ; la pipeline actuelle ne réalise pas de déploiement distant de FastAPI.",
          "Source : docs/deployment.md — protocole cible")
    steps(d, [("1. Construire", "Push et pull request\nTests automatiques\nArtefacts versionnés"),
              ("2. Vérifier", "Candidat en staging\nTests d’intégration\nValidation de version"),
              ("3. Mettre en service", "Déployer en production\nRedémarrer et vérifier\nSuivre les traces")])
    d.banner("Une mise à jour du dépôt Hugging Face ne redémarre pas l’API et ne réentraîne pas le modèle.", y=5.7)
    d.finish()

    d.new("GitHub Environments : variables, secrets et règles", "Déploiement · 3 / 4",
          "Le workflow sélectionne development pour les branches de travail, staging pour develop et production pour main, lors des push ou lancements manuels. Les pull requests exécutent les tests sans accéder à ces environnements. Chaque environnement GitHub peut contenir des variables et secrets propres et des règles de branches. Les approbations requises dépendent des fonctionnalités disponibles pour le dépôt. Le job environment-validation utilise une base PostgreSQL de test éphémère, vérifie le chemin du modèle puis relance les tests. Sur main, le job de publication utilise production pour envoyer le modèle vers SamGabb/hrnova. Il ne démarre pas de serveur API.",
          "Source : .github/workflows/ci.yml · docs/deployment.md")
    d.table(["Branche", "Environnement", "Rôle actuel dans HRNova"],
            [("Branche de travail", "development", "Validation de configuration et tests"),
             ("develop", "staging", "Validation de configuration et tests"),
             ("main", "production", "Validation puis publication Hugging Face")],
            .55, 2.0, [3.0, 2.7, 6.5], row_h=.8, font=17)
    d.banner("Un environnement GitHub protège un job ; il ne provisionne pas une API ni une base distante.", y=5.8)
    d.finish()

    d.new("Configurer, vérifier et pouvoir revenir en arrière", "Déploiement · 4 / 4",
          "DATABASE_URL doit pointer vers la base propre à la cible et HRNOVA_MODEL_PATH vers le fichier validé. Les fichiers .env.staging et .env.production ne sont pas versionnés. L’exemple suppose que la préproduction, ses tables, ses droits et son modèle ont déjà été préparés. Uvicorn écoute ici uniquement en local sur un port distinct. Avant de mettre en service une version, sauvegarder la base et conserver les références de l’ancienne version. Après démarrage, vérifier /health puis une inférence tracée. En cas d’incident, revenir au code et au modèle précédents ; une migration de schéma peut nécessiter un traitement supplémentaire. Aucun serveur de production n’est provisionné par cet exemple.",
          "Sources : docs/deployment.md · docs/maintenance.md")
    code(d, "# Exemple : instance staging déjà préparée\nuv run --no-sync uvicorn api:app \\\n  --env-file .env.staging \\\n  --host 127.0.0.1 --port 8002", .55, 1.95, 12.2, 1.9, 19)
    d.card(.55, 4.15, 5.95, 2.35, "Avant et après", "Sauvegarder → déployer → vérifier /health et une prédiction enregistrée en BDD.", size=17)
    d.card(6.8, 4.15, 5.95, 2.35, "Retour arrière", "Conserver les versions précédentes du code et du modèle, et prévoir les migrations.", size=17)
    d.finish()
    return d


def enrich(d):
    source = OUT / "HRNova_Soutenance_Enrichie_2026-10-29.pptx"
    if not source.is_file():
        return
    before = sha256(source.read_bytes()).hexdigest()
    prs = Presentation(source)
    assert len(prs.slides) == 32
    old = list(prs.slides._sldIdLst)
    original_texts = [tuple(s.text for s in slide.shapes if s.has_text_frame and not re.fullmatch(r"\d+\s*/\s*\d+", s.text.strip())) for slide in prs.slides]
    for src in d.prs.slides:
        dest = prs.slides.add_slide(prs.slide_layouts[6])
        for shape in list(dest.shapes):
            dest.shapes._spTree.remove(shape._element)
        for shape in src.shapes:
            dest.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
        dest.background.fill.solid()
        dest.background.fill.fore_color.rgb = src.background.fill.fore_color.rgb
        if dest.notes_slide.notes_text_frame is None:
            for shape in src.notes_slide.shapes:
                if shape.is_placeholder and shape._element.xpath('.//p:ph[@type="body"]'):
                    dest.notes_slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
        dest.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    added = list(prs.slides._sldIdLst)[32:]
    # Après la section Git/CI, avant la documentation et les annexes.
    order = old[:27] + added + old[27:]
    for item in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(item)
    for item in order:
        prs.slides._sldIdLst.append(item)
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.left > Inches(10) and re.fullmatch(r"\d+\s*/\s*\d+", shape.text.strip()):
                runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
                if runs:
                    runs[0].text = f"{i:02d} / 36"
                    for run in runs[1:]:
                        run.text = ""
    target = OUT / "HRNova_Soutenance_Environnements_2026-10-29.pptx"
    prs.save(target)
    check = Presentation(target)
    assert len(check.slides) == 36
    for texts, position in zip(original_texts, [i for i, item in enumerate(order) if item in old]):
        assert texts == tuple(s.text for s in check.slides[position].shapes if s.has_text_frame and not re.fullmatch(r"\d+\s*/\s*\d+", s.text.strip()))
    assert before == sha256(source.read_bytes()).hexdigest()
    print(f"Soutenance : {target.name} — 36 diapositives, source conservée")


def main():
    d = build()
    assert len(d.prs.slides) == 4
    OUT.mkdir(exist_ok=True)
    d.prs.save(OUT / f"{STEM}.pptx")
    overview = Image.new("RGB", (1600, 900), "white")
    notes = ["# Déploiement par environnement — Notes pour l’oral", ""]
    for i, (image, (title, note, source)) in enumerate(zip(d.images, d.notes)):
        overview.paste(image.resize((800, 450)), ((i % 2) * 800, (i // 2) * 450))
        notes.extend([f"## {i + 1}. {title}", "", note, "", source, ""])
    overview.save(OUT / f"{STEM}_Apercu.png")
    (OUT / f"{STEM}_Notes.md").write_text("\n".join(notes), encoding="utf-8")
    for slide in d.prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= d.prs.slide_width + 10
            assert shape.top + shape.height <= d.prs.slide_height + 10
    enrich(d)
    print("Support déploiement : 4 diapositives avec notes et aperçu")


if __name__ == "__main__":
    main()
